"""Small helpers for reading/writing the Excel files."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

ALIASES = {
    "email": ("email", "e-mail", "email address", "mail"),
    "name": ("name", "contact", "contact name", "first name", "owner"),
    "company": ("company", "business", "business name", "company name"),
    "website": ("website", "url", "site", "web", "web address"),
}

DRAFT_HEADERS = ["email", "name", "company", "website", "subject", "body", "status", "notes", "sent_at"]
STATUSES = ["review", "approved", "skip", "sent", "failed", "unsubscribed"]
COLUMN_WIDTHS = {"email": 28, "name": 16, "company": 22, "website": 28, "subject": 36,
                 "body": 80, "status": 12, "notes": 36, "sent_at": 18}


class FileLocked(Exception):
    pass


def read_rows(path):
    try:
        wb = load_workbook(path, data_only=True)
    except PermissionError as exc:
        raise FileLocked(f"Cannot open {path}. Close it in Excel and try again.") from exc
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in r):
            continue
        out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(headers, r) if h})
    return out


def pick(row, field):
    """Read a field from a lead row, accepting common column-name variants."""
    for alias in ALIASES[field]:
        if row.get(alias):
            return row[alias]
    return ""


def write_drafts(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "drafts"
    ws.append(DRAFT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h, "") for h in DRAFT_HEADERS])
    for idx, header in enumerate(DRAFT_HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = COLUMN_WIDTHS[header]
    for line in ws.iter_rows(min_row=2):
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    status_col = ws.cell(row=1, column=DRAFT_HEADERS.index("status") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{max(len(rows) + 1, 2)}")
    try:
        wb.save(path)
    except PermissionError as exc:
        raise FileLocked(f"Cannot save {path}. Close it in Excel and run the command again "
                         "(progress is kept).") from exc
